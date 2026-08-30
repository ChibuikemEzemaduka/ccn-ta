"""
generate_variants.py  —  practice-problem generator (Prof. Kar's suggestion)

Give it a homework file and its official solutions file. For each problem it
generates a NEW variant -- same concept, same method, different numbers or
scenario -- and a freshly worked solution for that new variant, derived from
the original solution's method.

Generation is deliberately split into two separate model calls per problem:
  1. Vary the problem (numbers/scenario changed, same concept and difficulty).
  2. Solve the NEW problem from scratch, low-temperature, with a self-check
     step -- kept separate from step 1 so creative variation doesn't also
     have to be the thing doing careful arithmetic.

IMPORTANT: these are AI-generated. They are a strong first draft, not a
guarantee. Review every problem and solution before handing them to students
-- this warning is also embedded in the exported documents themselves.

Usage:
    python generate_variants.py path\\to\\HW4.pdf path\\to\\HW4_solutions.pdf

Produces, in a 'generated' folder:
    HW4_practice_problems.docx     (questions only -- safe to hand to students)
    HW4_practice_answer_key.docx   (questions + solutions -- instructor only)
    HW4_variants.json              (raw data, for the dashboard / re-export)
"""

import sys
import io
import json
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

import add_homework   # reuses read_text, llm_split, regex_split
import llm_client   # ****** see llm_client.py to change which model/provider is used ******

# **** ALIAS, NOT THE SOURCE OF TRUTH -- edit llm_client.py to change models ****
MODEL = llm_client.TUTOR_MODEL
OUTPUT_DIR = Path("generated")

VARY_INSTRUCTION = """You create a NEW practice variant of a homework problem for an instructor.

You will see the ORIGINAL PROBLEM and the FULL OFFICIAL SOLUTIONS DOCUMENT for the
homework set (the relevant solution may be anywhere in it -- use it to understand
the METHOD being tested, independent of the specific numbers). The original
problem may ALSO be provided to you as page images that can contain diagrams,
figures, network topologies, or tables -- read those images carefully and take
any visual content into account when creating the variant.

Create a new problem that:
- Tests the EXACT same concept and requires the EXACT same method/formula/steps.
- Is NOT identical: change the numbers, names, or surface scenario (a different
  prefix length, host count, link speed, window size, protocol parameter, etc.).
- For "explain"/"prove"/"derive" problems with no numbers, vary the specific
  scenario or setup instead, so it is not a verbatim repeat of the original.
- Remains realistic, solvable, and well-posed with the same approach as the original.
- Is self-contained: a student should be able to read ONLY your new problem
  text and solve it, without needing the original.

If the problem REQUIRES a figure to be well-posed (a network topology, a routing
table, a state machine, a timing diagram, etc.), provide TWO things:
 1. "figure_description": a clear human-readable description of the figure.
 2. "figure_spec": a STRUCTURED, renderable specification the system can draw
    into an actual image. Use this JSON shape when a node-link diagram fits
    (network topology, state machine, graph):
      {"type": "graph",
       "directed": true/false,
       "nodes": [{"id": "A", "label": "Router A"}, ...],
       "edges": [{"from": "A", "to": "B", "label": "10 Mbps"}, ...]}
    Or for a table (routing table, address table):
      {"type": "table",
       "headers": ["Destination", "Next Hop", "Cost"],
       "rows": [["10.0.0.0/8", "R2", "3"], ...]}
    If no figure is needed, set BOTH "figure_description" and "figure_spec" to
    empty (""/null). Only include a figure when the problem genuinely needs one.

Respond with ONLY a JSON object:
{"variant_problem": "<the new, self-contained problem statement>",
 "figure_description": "<human-readable description of any figure, or empty>",
 "figure_spec": <structured figure object as above, or null>,
 "what_changed": "<one sentence: what you changed from the original>"}"""

SOLVE_INSTRUCTION = """You solve a homework problem for an instructor's answer key.
Accuracy matters more than anything else here -- this will be used to check
student work, so an arithmetic mistake is a real problem.

You will see: the ORIGINAL problem, the ORIGINAL official solution (showing the
method to follow), and a NEW variant problem that uses the same method with
different numbers or scenario.

Solve the NEW variant problem by carrying the original method through with its
own numbers/scenario. Show your work step by step, the way the original solution
does. Before finalizing, recompute the final result once more to check it is
internally consistent with your own steps.

Respond with ONLY a JSON object:
{"variant_solution": "<step-by-step worked solution, ending in the final result>"}"""


def _json_call(client, instruction, prompt, temperature, attachment_parts=None):
    turn = {"role": "user", "text": prompt}
    if attachment_parts:
        turn["attachments"] = attachment_parts
    reply_text = client.generate(MODEL, instruction, [turn],
                                 temperature=temperature, json_mode=True)
    return json.loads(reply_text)


def _vary_instruction(instructor_prompt: str = "") -> str:
    """Return VARY_INSTRUCTION. When the instructor gave specific instructions,
    put them FIRST and emphatically -- models weight the top of the system prompt
    most heavily, so leading with the instructor's intent is what makes the
    variant faithfully follow it (appending at the end gets ignored more often)."""
    if not instructor_prompt.strip():
        return VARY_INSTRUCTION
    return (
        "########################################################\n"
        "# TOP PRIORITY -- THE INSTRUCTOR'S INSTRUCTION FOR THIS VARIANT\n"
        "########################################################\n"
        f"{instructor_prompt.strip()}\n\n"
        "You MUST follow the instruction above. It OVERRIDES any conflicting "
        "default below. If it says to increase difficulty, make it harder than "
        "the original. If it says to focus on a specific concept, center the "
        "variant on that concept. If it names a specific change (topology, "
        "numbers, protocol, addressing scheme), apply exactly that change. Do "
        "not ignore or water down the instruction.\n"
        "########################################################\n\n"
        + VARY_INSTRUCTION
    )


def select_target_problems(instructor_prompt, problems):
    """Given the instructor's instruction and the list of (number, text) problems,
    return the subset the instructor wants variants for.

    Detects phrasings like 'only problem 2', 'just P3', 'problems 1 and 4',
    'generate a variant for question 5', 'problem 2 only'. If no specific problem
    is named, returns ALL problems (the default -- generate for everything).

    This is a best-effort text match; when in doubt it includes the problem, so
    the instructor never silently loses a problem they wanted."""
    if not instructor_prompt or not instructor_prompt.strip():
        return problems

    import re
    text = instructor_prompt.lower()

    # Only treat it as targeting if the instruction actually scopes to specific
    # problems. Look for problem/question/P<n> references.
    # Collect all referenced numbers.
    nums = set()
    # "problem 2", "question 3", "part 4"
    for m in re.finditer(r'(?:problem|question|part|prob|q|p)\s*#?\s*(\d+)', text):
        nums.add(int(m.group(1)))
    # "P2", "P3" tokens
    for m in re.finditer(r'\bp(\d+)\b', text):
        nums.add(int(m.group(1)))

    # Was the instruction actually SCOPING (only/just/specifically), or merely
    # mentioning a number in passing? Require a scoping cue OR that the only
    # content is the problem reference.
    scoping = any(w in text for w in ("only", "just", "specifically", "nothing else",
                                       "solely", "exclusively"))
    if nums and (scoping or re.search(r'^\s*(?:generate|make|create|do)?\s*'
                                      r'(?:a\s+)?variant', text)):
        selected = [(n, t) for (n, t) in problems if n in nums]
        if selected:
            return selected
    # No specific targeting detected -> generate for all problems.
    return problems


def render_figure_spec(figure_spec):
    """Render a structured figure_spec into an actual PNG image (bytes), or
    return None if it can't be rendered. Supports node-link graphs (topologies,
    state machines) and tables (routing tables). This is how the module produces
    its OWN images for variants, not just text descriptions.

    Best-effort: if the needed libraries aren't available or the spec is
    malformed, returns None and the caller falls back to the text description."""
    if not figure_spec or not isinstance(figure_spec, dict):
        return None
    kind = figure_spec.get("type")
    try:
        import matplotlib
        matplotlib.use("Agg")   # headless backend
        import matplotlib.pyplot as plt
        import io as _io
    except ImportError:
        return None

    try:
        if kind == "graph":
            return _render_graph(figure_spec, plt, _io)
        if kind == "table":
            return _render_table(figure_spec, plt, _io)
    except Exception:
        return None
    return None


def _render_graph(spec, plt, _io):
    """Render a node-link diagram (network topology / state machine)."""
    try:
        import networkx as nx
    except ImportError:
        return None
    directed = bool(spec.get("directed"))
    G = nx.DiGraph() if directed else nx.Graph()
    labels = {}
    for n in spec.get("nodes", []):
        nid = str(n.get("id"))
        G.add_node(nid)
        labels[nid] = n.get("label", nid)
    edge_labels = {}
    for e in spec.get("edges", []):
        a, b = str(e.get("from")), str(e.get("to"))
        G.add_edge(a, b)
        if e.get("label"):
            edge_labels[(a, b)] = e["label"]
    if G.number_of_nodes() == 0:
        return None
    fig, ax = plt.subplots(figsize=(6, 4.2))
    pos = nx.spring_layout(G, seed=42)   # deterministic layout
    nx.draw_networkx_nodes(G, pos, ax=ax, node_color="#E8EEF7",
                           edgecolors="#2563EB", node_size=1600)
    nx.draw_networkx_labels(G, pos, labels=labels, ax=ax, font_size=10)
    if directed:
        nx.draw_networkx_edges(G, pos, ax=ax, edge_color="#5B6470",
                               arrows=True, arrowsize=18, width=1.6)
    else:
        nx.draw_networkx_edges(G, pos, ax=ax, edge_color="#5B6470", width=1.6)
    if edge_labels:
        nx.draw_networkx_edge_labels(G, pos, edge_labels=edge_labels, ax=ax,
                                     font_size=8, font_color="#C8102E")
    ax.axis("off")
    fig.tight_layout()
    buf = _io.BytesIO()
    fig.savefig(buf, format="png", dpi=150, bbox_inches="tight")
    plt.close(fig)
    return buf.getvalue()


def _render_table(spec, plt, _io):
    """Render a table (routing table, address table) as an image."""
    headers = spec.get("headers", [])
    rows = spec.get("rows", [])
    if not headers or not rows:
        return None
    fig, ax = plt.subplots(figsize=(min(8, 1.6 * len(headers)),
                                    0.5 + 0.4 * (len(rows) + 1)))
    ax.axis("off")
    table = ax.table(cellText=rows, colLabels=headers, loc="center",
                     cellLoc="center")
    table.auto_set_font_size(False)
    table.set_fontsize(10)
    table.scale(1, 1.4)
    # Header styling
    for j in range(len(headers)):
        cell = table[0, j]
        cell.set_facecolor("#C8102E")
        cell.set_text_props(color="white", fontweight="bold")
    fig.tight_layout()
    buf = _io.BytesIO()
    fig.savefig(buf, format="png", dpi=150, bbox_inches="tight")
    plt.close(fig)
    return buf.getvalue()


def generate_one_variant(client, problem_id, original_problem, solutions_text,
                         instructor_prompt: str = "", attachment_parts=None):
    """Returns a dict with the variant problem+solution, or {'error': ...} on failure.

    instructor_prompt: free-text instructions from the instructor chat that
    direct how this specific variant should be generated (e.g. 'increase the
    difficulty', 'focus on CIDR notation', 'change the topology to a ring').
    When provided it is injected into the VARY step; the SOLVE step always
    follows the original method regardless.

    attachment_parts: optional list of image attachments (page images of the
    uploaded problem document) so the model can SEE diagrams, figures, and
    tables embedded in the original problem when creating the variant."""
    try:
        step1 = _json_call(
            client, _vary_instruction(instructor_prompt),
            f"ORIGINAL PROBLEM [{problem_id}]:\n{original_problem}\n\n"
            f"FULL OFFICIAL SOLUTIONS DOCUMENT:\n{solutions_text}",
            temperature=0.6,
            attachment_parts=attachment_parts,
        )
        variant_problem = step1["variant_problem"]
        what_changed = step1.get("what_changed", "")
        figure_description = step1.get("figure_description", "") or ""
        figure_spec = step1.get("figure_spec")

        # If the variant needs a figure, (a) render it to an actual image, and
        # (b) append its text spec to the problem so it also travels in exports.
        figure_image = render_figure_spec(figure_spec)
        if figure_description.strip():
            variant_problem = (variant_problem
                               + "\n\n[Figure for this problem]\n" + figure_description.strip())

        step2 = _json_call(
            client, SOLVE_INSTRUCTION,
            f"ORIGINAL PROBLEM:\n{original_problem}\n\n"
            f"ORIGINAL OFFICIAL SOLUTIONS DOCUMENT (find the relevant method in here):\n{solutions_text}\n\n"
            f"NEW VARIANT PROBLEM TO SOLVE:\n{variant_problem}",
            temperature=0.0,
        )
        variant_solution = step2["variant_solution"]

        return {
            "problem_id": problem_id,
            "original_problem": original_problem,
            "variant_problem": variant_problem,
            "variant_solution": variant_solution,
            "figure_description": figure_description,
            "figure_spec": figure_spec,
            "figure_image": figure_image,   # PNG bytes, or None
            "what_changed": what_changed,
            "instructor_prompt": instructor_prompt,
            "error": None,
        }
    except Exception as error:
        return {
            "problem_id": problem_id,
            "original_problem": original_problem,
            "variant_problem": None,
            "variant_solution": None,
            "figure_description": None,
            "figure_spec": None,
            "figure_image": None,
            "what_changed": None,
            "instructor_prompt": instructor_prompt,
            "error": str(error)[:200],
        }


def generate_variant_set(client, homework_text, solutions_text,
                         instructor_prompt: str = "", attachment_parts=None):
    """Split the homework into problems and generate one variant per problem.
    Returns a list of variant dicts (see generate_one_variant).

    instructor_prompt is forwarded to every individual generate_one_variant
    call so all variants in the set follow the same instruction.
    attachment_parts (page images of the uploaded problem doc) are forwarded so
    embedded diagrams/figures are understood."""
    problems = add_homework.llm_split(homework_text, client) if client else None
    if problems is None:
        problems = add_homework.regex_split(homework_text)
    if problems is None:
        problems = [(1, homework_text.strip())]

    # Honour any 'only problem N' targeting in the instruction.
    problems = select_target_problems(instructor_prompt, problems)

    variants = []
    for number, text in problems:
        problem_id = f"P{number}"
        variants.append(generate_one_variant(client, problem_id, text, solutions_text,
                                             instructor_prompt=instructor_prompt,
                                             attachment_parts=attachment_parts))
    return variants


# ------------------------------ Word export --------------------------------

DISCLAIMER = ("\u26a0 AI-GENERATED. Review every problem and solution carefully before "
             "sharing with students -- check the math and the wording.")


def _add_disclaimer(doc):
    p = doc.add_paragraph()
    run = p.add_run(DISCLAIMER)
    run.bold = True
    run.font.color.rgb = RGBColor(0xB0, 0x00, 0x00)
    run.font.size = Pt(10)


def _add_figure(doc, v):
    """Embed the generated figure image into the doc, if there is one."""
    img = v.get("figure_image")
    if not img:
        return
    try:
        from docx.shared import Inches
        doc.add_picture(io.BytesIO(img), width=Inches(4.5))
    except Exception:
        pass


def build_questions_docx(hw_id, variants):
    """Questions only -- safe to hand directly to students."""
    doc = Document()
    doc.add_heading(f"{hw_id} -- Practice Variant Set", level=1)
    _add_disclaimer(doc)
    for v in variants:
        if v.get("error"):
            continue
        doc.add_heading(f"Problem {v['problem_id']}", level=2)
        doc.add_paragraph(v["variant_problem"])
        _add_figure(doc, v)
    return doc


def build_answer_key_docx(hw_id, variants):
    """Questions + solutions -- instructor only."""
    doc = Document()
    doc.add_heading(f"{hw_id} -- Practice Variant Set: Answer Key", level=1)
    _add_disclaimer(doc)
    for v in variants:
        doc.add_heading(f"Problem {v['problem_id']}", level=2)
        if v.get("error"):
            doc.add_paragraph(f"(Generation failed for this problem: {v['error']})")
            continue
        doc.add_heading("Problem", level=3)
        doc.add_paragraph(v["variant_problem"])
        _add_figure(doc, v)
        doc.add_heading("Solution", level=3)
        doc.add_paragraph(v["variant_solution"])
        if v.get("what_changed"):
            note = doc.add_paragraph()
            note.add_run(f"Changed from original: {v['what_changed']}").italic = True
    return doc


def doc_to_bytes(doc):
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# -------------------------------- CLI ---------------------------------------

def main():
    if len(sys.argv) != 3:
        print("Usage: python generate_variants.py <homework_file> <solutions_file>")
        return
    hw_path, sol_path = Path(sys.argv[1]), Path(sys.argv[2])
    if not hw_path.exists() or not sol_path.exists():
        print("One of those files doesn't exist. Check the paths.")
        return

    ok, info = llm_client.check_connection()
    if not ok:
        print(f"Can't reach the local model server (Ollama) at {llm_client.OLLAMA_BASE_URL} -- "
              f"this tool needs it to generate variants. ({info})")
        return
    client = llm_client

    hw_id = add_homework.hw_id_from_filename(hw_path.name)
    homework_text = add_homework.read_text(hw_path)
    solutions_text = add_homework.read_text(sol_path)

    print(f"Splitting {hw_path.name} into problems...")
    variants = generate_variant_set(client, homework_text, solutions_text)
    ok = sum(1 for v in variants if not v.get("error"))
    print(f"Generated {ok}/{len(variants)} variant(s).")
    for v in variants:
        status = "OK" if not v.get("error") else f"FAILED ({v['error']})"
        print(f"  {v['problem_id']}: {status}")

    OUTPUT_DIR.mkdir(exist_ok=True)
    json_path = OUTPUT_DIR / f"{hw_id}_variants.json"
    json_path.write_text(json.dumps(variants, indent=2, ensure_ascii=False), encoding="utf-8")

    build_questions_docx(hw_id, variants).save(OUTPUT_DIR / f"{hw_id}_practice_problems.docx")
    build_answer_key_docx(hw_id, variants).save(OUTPUT_DIR / f"{hw_id}_practice_answer_key.docx")

    print(f"\nSaved to {OUTPUT_DIR}/:")
    print(f"  {hw_id}_practice_problems.docx    (hand-out safe)")
    print(f"  {hw_id}_practice_answer_key.docx  (instructor only)")
    print(f"  {hw_id}_variants.json")
    print(f"\n{DISCLAIMER}")


if __name__ == "__main__":
    main()


# ─────────────────────────── Wireshark question generation ───────────────────

WIRESHARK_Q_INSTRUCTION = """You generate Wireshark-trace based homework questions for a Computer Communication Networks course.

You receive a parsed Wireshark capture summary (packet list with timestamps, IPs, protocols, flags, sequence numbers, etc.) and a topic the instructor wants to test.

For each question you generate:
- A self-contained QUESTION a student can answer by examining the trace.
- A precise ANSWER (cite exact packet numbers, field values, or computed quantities from the trace -- do not be vague).
- 3 SOCRATIC HINTS given in order: each hint moves the student one step closer to the answer without revealing it. Hint 1 points to the right section of the trace; Hint 2 names the specific field/column to look at; Hint 3 gives a partial computation or direction.
- PACKETS_REFERENCED: a short list of packet numbers the question is about.

Rules:
- Questions must be answerable from THIS trace only (not from memory or general knowledge alone).
- Vary difficulty across the set: include at least one observation question (read a value) and one analysis question (infer or compute something).
- Do not repeat essentially the same question.
- If the trace doesn't have enough content for the requested number of questions at the requested topic, generate as many distinct ones as you can and note any shortfall.

Respond ONLY with a JSON array:
[
  {
    "question": "...",
    "answer": "...",
    "hints": ["hint 1 ...", "hint 2 ...", "hint 3 ..."],
    "packets_referenced": "e.g. 1-3, 7"
  },
  ...
]"""


def generate_wireshark_questions(client, trace_summary, topic, n_questions=5):
    """Generate n_questions Wireshark-trace-based questions with Socratic hints.
    Returns a list of dicts: {question, answer, hints, packets_referenced, error}."""
    prompt = (f"TOPIC TO TEST: {topic}\n\n"
              f"NUMBER OF QUESTIONS: {n_questions}\n\n"
              f"WIRESHARK CAPTURE SUMMARY:\n{trace_summary}")
    try:
        reply = client.generate(
            llm_client.TUTOR_MODEL,
            WIRESHARK_Q_INSTRUCTION,
            [{"role": "user", "text": prompt}],
            temperature=0.4, json_mode=True,
        )
        items = json.loads(reply)
        if not isinstance(items, list):
            return [{"question": "", "answer": "", "hints": [], "error": "Unexpected response format"}]
        result = []
        for i, item in enumerate(items[:n_questions]):
            result.append({
                "question": item.get("question", ""),
                "answer": item.get("answer", ""),
                "hints": item.get("hints", []),
                "packets_referenced": item.get("packets_referenced", ""),
                "error": None,
            })
        return result
    except Exception as error:
        return [{"question": "", "answer": "", "hints": [], "packets_referenced": "",
                 "error": str(error)[:200]}]


def build_wireshark_questions_docx(trace_name, topic, questions):
    """Student-facing question sheet (no answers, no hints)."""
    doc = Document()
    doc.add_heading(f"Wireshark Exercise: {trace_name}", level=1)
    doc.add_heading(f"Topic: {topic}", level=2)
    _add_disclaimer(doc)
    ok_qs = [q for q in questions if not q.get("error")]
    for i, q in enumerate(ok_qs, 1):
        doc.add_heading(f"Question {i}", level=3)
        if q.get("packets_referenced"):
            p = doc.add_paragraph()
            p.add_run(f"Refer to: packets {q['packets_referenced']}").italic = True
        doc.add_paragraph(q["question"])
    return doc


def build_wireshark_answer_key_docx(trace_name, topic, questions):
    """Instructor answer key with answers and Socratic hints."""
    doc = Document()
    doc.add_heading(f"Wireshark Exercise Answer Key: {trace_name}", level=1)
    doc.add_heading(f"Topic: {topic}", level=2)
    _add_disclaimer(doc)
    for i, q in enumerate(questions, 1):
        doc.add_heading(f"Question {i}", level=3)
        if q.get("error"):
            doc.add_paragraph(f"(Generation failed: {q['error']})")
            continue
        if q.get("packets_referenced"):
            p = doc.add_paragraph()
            p.add_run(f"Refers to: packets {q['packets_referenced']}").italic = True
        doc.add_heading("Question", level=4)
        doc.add_paragraph(q["question"])
        doc.add_heading("Answer", level=4)
        doc.add_paragraph(q["answer"])
        if q.get("hints"):
            doc.add_heading("Socratic Hints (give one at a time)", level=4)
            for j, hint in enumerate(q["hints"], 1):
                doc.add_paragraph(f"Hint {j}: {hint}")
    return doc
